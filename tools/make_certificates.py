"""Make award certificates (with Certificate No. + verification QR) for the top 3 of a marksheet.

Usage:
  python3 tools/make_certificates.py MARKSHEET.xlsx "Eng Oct 2026" [--issued 2026-11-01] [--out certificates_out]

- Reads the TOP 3 box of that month's sheet (filled by Excel when the marksheet was saved).
- Adds each winner to site/data/certificates.json (used by site/verify.html).
- Writes a Word file with one certificate per winner, each with its own QR code.
Use --blank to make the blank template instead:  python3 tools/make_certificates.py --blank
"""
import argparse, datetime, json, os, re, subprocess, sys
import qrcode
from qrcode.constants import ERROR_CORRECT_M

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB = os.path.join(ROOT, "site", "data", "certificates.json")
VERIFY = "https://altafhussainbirhmani376-rgb.github.io/TalentCore-Global/verify.html"
INSTITUTE = "Future Talent Academy Jamshoro"
MONTHS = {m[:3]: m for m in ["January", "February", "March", "April", "May", "June", "July", "August",
                              "September", "October", "November", "December"]}
CLASSES = {"Eng": ("English", "ENG"), "Com": ("Computer", "COM")}


def qr_png(url, path):
    q = qrcode.QRCode(error_correction=ERROR_CORRECT_M, box_size=10, border=1)
    q.add_data(url)
    q.make(fit=True)
    q.make_image(fill_color="#0D3B66", back_color="white").save(path)


def winners_from_sheet(xlsx, sheet):
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, data_only=True)[sheet]
    rows = {}
    for r in range(13, 63):  # student rows: B name, C father, I percentage, K position
        name = ws[f"B{r}"].value
        if name:
            rows[name] = (ws[f"C{r}"].value or "", ws[f"I{r}"].value)
    out = []
    for r in range(8, 12):  # TOP 3 box
        pos, name = ws[f"A{r}"].value, ws[f"B{r}"].value
        if pos in (None, "") or not name:
            continue
        father, pct = rows.get(name, ("", None))
        out.append({"position": int(pos), "name": str(name), "father": str(father),
                    "percentage": f"{pct * 100:.1f}%" if isinstance(pct, (int, float)) else ""})
    if not out:
        sys.exit(f"No winners found on sheet '{sheet}'. Open the marksheet in Excel, check the marks, save, then try again.")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?")
    ap.add_argument("sheet", nargs="?")
    ap.add_argument("--issued", default=datetime.date.today().isoformat())
    ap.add_argument("--out", default=os.path.join(ROOT, "certificates_out"))
    ap.add_argument("--blank", action="store_true")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    tmp = os.path.join(a.out, "_qr")
    os.makedirs(tmp, exist_ok=True)

    if a.blank:
        qr_png(VERIFY, os.path.join(tmp, "blank.png"))
        spec = {"out": os.path.join(a.out, "Award_Certificates_Blank_Template.docx"),
                "certs": [{"position": p, "blank": True, "qr": os.path.join(tmp, "blank.png")} for p in (1, 2, 3)]}
    else:
        if not (a.xlsx and a.sheet):
            ap.error("give the marksheet file and the sheet name, or use --blank")
        m = re.match(r"(Eng|Com)\s+([A-Za-z]{3})\s+(\d{4})$", a.sheet)
        if not m:
            sys.exit('Sheet name must look like "Eng Oct 2026" or "Com Oct 2026".')
        cls, code = CLASSES[m.group(1)]
        month = f"{MONTHS[m.group(2).title()]} {m.group(3)}"
        yymm = f"{m.group(3)[2:]}{list(MONTHS).index(m.group(2).title()) + 1:02d}"
        db = json.load(open(DB, encoding="utf-8")) if os.path.exists(DB) else []
        db = [c for c in db if not c["id"].startswith(f"TCG-FTA-{code}-{yymm}-")]  # re-running replaces that month
        certs = []
        for i, w in enumerate(winners_from_sheet(a.xlsx, a.sheet), 1):
            cid = f"TCG-FTA-{code}-{yymm}-{i:02d}"
            rec = {"id": cid, "name": w["name"], "father": w["father"], "class": f"{cls} Class",
                   "position": w["position"], "month": month, "percentage": w["percentage"],
                   "institute": INSTITUTE, "issued": a.issued}
            db.append(rec)
            png = os.path.join(tmp, cid + ".png")
            qr_png(f"{VERIFY}?id={cid}", png)
            certs.append({**rec, "qr": png})
            print(f"{cid}  {w['position']}  {w['name']}  {w['percentage']}")
        db.sort(key=lambda c: c["id"])
        with open(DB, "w", encoding="utf-8") as f:
            json.dump(db, f, ensure_ascii=False, indent=2)
            f.write("\n")
        spec = {"out": os.path.join(a.out, f"Certificates_{cls}_{month.replace(' ', '_')}.docx"), "certs": certs}

    spec_path = os.path.join(tmp, "spec.json")
    json.dump(spec, open(spec_path, "w", encoding="utf-8"), ensure_ascii=False)
    subprocess.run(["node", os.path.join(ROOT, "tools", "certificate_docx.js"), spec_path], check=True)
    print("Saved:", spec["out"])


if __name__ == "__main__":
    main()
